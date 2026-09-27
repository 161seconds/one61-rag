"""Document text extraction — PDF, DOCX, TXT.

Extracts plain text from uploaded files for RAG chunking.
Gracefully handles missing optional dependencies.
"""

import io
import logging

logger = logging.getLogger(__name__)

# Max file size: 25MB
MAX_FILE_SIZE = 25 * 1024 * 1024


def extract_text(content: bytes, filename: str) -> str:
    """Extract plain text from file content based on file extension.

    Args:
        content: Raw file bytes.
        filename: Original filename (used to detect format).

    Returns:
        Extracted plain text string.

    Raises:
        ValueError: If file is too large or content is empty.
        ImportError: If required parser library is not installed.
    """
    if not content:
        raise ValueError("Empty file content")

    if len(content) > MAX_FILE_SIZE:
        raise ValueError(
            f"File too large: {len(content) / 1024 / 1024:.1f}MB "
            f"(max {MAX_FILE_SIZE / 1024 / 1024:.0f}MB)"
        )

    ext = _get_extension(filename)
    logger.info("Extracting text from '%s' (ext=%s, size=%dKB)", filename, ext, len(content) // 1024)

    if ext == "pdf":
        return _extract_pdf(content)
    elif ext == "docx":
        return _extract_docx(content)
    elif ext in ("txt", "md", "csv", "json", "xml", "html", "htm"):
        return _extract_text(content)
    else:
        # Fallback: attempt plain text decode
        logger.warning("Unknown extension '%s', attempting plain text decode", ext)
        return _extract_text(content)


def _get_extension(filename: str) -> str:
    """Extract lowercase file extension from filename."""
    if "." not in filename:
        return ""
    return filename.rsplit(".", 1)[-1].lower()


def _extract_text(content: bytes) -> str:
    """Decode as UTF-8 plain text."""
    return content.decode("utf-8", errors="replace").strip()


def _extract_pdf(content: bytes) -> str:
    """Extract text from PDF using PyMuPDF (fitz)."""
    try:
        import fitz  # pymupdf
    except ImportError:
        raise ImportError(
            "pymupdf is required for PDF extraction. "
            "Install with: pip install pymupdf"
        )

    doc = fitz.open(stream=content, filetype="pdf")
    text_parts = []
    page_count = doc.page_count
    for page in doc:
        page_text = page.get_text("text")
        if page_text.strip():
            text_parts.append(page_text.strip())
    doc.close()

    result = "\n\n".join(text_parts)
    logger.info("PDF extraction: %d pages, %d chars", page_count, len(result))
    return result


def _extract_docx(content: bytes) -> str:
    """Extract text from DOCX using python-docx."""
    try:
        from docx import Document
    except ImportError:
        raise ImportError(
            "python-docx is required for DOCX extraction. "
            "Install with: pip install python-docx"
        )

    doc = Document(io.BytesIO(content))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    result = "\n\n".join(paragraphs)
    logger.info("DOCX extraction: %d paragraphs, %d chars", len(paragraphs), len(result))
    return result
